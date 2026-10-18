"""Generate synthetic PDF and DOCX fixtures in memory (no binary files in the repo)."""

import io


def _escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(pages: list[list[str]]) -> bytes:
    """A well-formed PDF with one page per entry; each string is one line of text.

    Pages given as an empty list have no text layer (like a scanned page).
    """
    page_count = len(pages)
    # Object numbers: 1 catalog, 2 pages, 3 font, then (page, content) pairs.
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids = []
    for index, lines in enumerate(pages):
        page_number = 4 + index * 2
        content_number = page_number + 1
        kids.append(f"{page_number} 0 R")
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 740 Td"]
        for line in lines:
            ops.append(f"({_escape_pdf_text(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects[page_number] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_number
        )
        objects[content_number] = (
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {page_count} >>".encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"
    size = max(objects) + 1
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for number in range(1, size):
        out += b"%010d 00000 n \n" % offsets[number]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref)
    return bytes(out)


def build_encrypted_pdf(text: str, password: str = "fixture-password") -> bytes:  # noqa: S107
    from pypdf import PdfReader, PdfWriter  # noqa: PLC0415

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(build_pdf([[text]]))))
    writer.encrypt(user_password=password, owner_password=password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx  # noqa: PLC0415

    document = docx.Document()
    for index, paragraph in enumerate(paragraphs):
        if index == 0:
            document.add_heading(paragraph, level=1)
        else:
            document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                grid.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
